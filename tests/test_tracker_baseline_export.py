from __future__ import annotations

from datetime import datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from src.gui.tracker_baseline_compare import BaselineComparisonSource
from src.gui.tracker_baseline_compare import compare_tracker_items
from src.gui.tracker_baseline_export import BaselineExportError
from src.gui.tracker_baseline_export import baseline_export_fields
from src.gui.tracker_baseline_export import create_baseline_comparison_workbook
from src.gui.tracker_baseline_export import export_baseline_comparison_xlsx
from src.gui.tracker_query_models import TrackerItemSummary


def _item(
    item_id: int,
    *,
    name: str,
    status: str,
    assignees: list[dict] | None = None,
    table_rows: list[list[dict]] | None = None,
) -> TrackerItemSummary:
    return TrackerItemSummary.from_raw(
        {
            "id": item_id,
            "name": name,
            "status": {
                "id": 1 if status == "Open" else 2,
                "name": status,
                "type": "ChoiceOptionReference",
            },
            "assignedTo": assignees or [],
            "customFields": [
                {
                    "fieldId": 10,
                    "name": "검증 표",
                    "type": "TableFieldValue",
                    "values": table_rows or [],
                }
            ],
        },
        tracker_id=20,
    )


def _table_row(condition: str, result: bool) -> list[dict]:
    return [
        {
            "fieldId": 501,
            "name": "조건",
            "type": "TextFieldValue",
            "value": condition,
        },
        {
            "fieldId": 502,
            "name": "결과",
            "type": "BoolFieldValue",
            "value": result,
        },
    ]


class TrackerBaselineExportTest(unittest.TestCase):
    def setUp(self) -> None:
        comparison = _item(
            1,
            name="이전 로그인 기능",
            status="Draft",
            assignees=[{"id": 7, "name": "홍길동", "type": "UserReference"}],
            table_rows=[_table_row("로그인", True), _table_row("권한 확인", False)],
        )
        reference = _item(
            1,
            name="로그인 기능",
            status="Open",
            assignees=[
                {"id": 7, "name": "홍길동", "type": "UserReference"},
                {"id": 9, "name": "김영희", "type": "UserReference"},
            ],
            table_rows=[_table_row("로그인", True), _table_row("권한 확인", True)],
        )
        added = _item(2, name="비밀번호 재설정", status="Open")
        removed = _item(3, name="기존 로그인", status="Draft")
        self.result = compare_tracker_items(
            (comparison, removed),
            (reference, added),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )
        self.selected_keys = ("name", "status", "assignedTo", "custom:10")

    def test_fields_include_table_metadata_and_column_union(self) -> None:
        fields = {field.field_key: field for field in baseline_export_fields(self.result)}

        self.assertTrue(fields["custom:10"].is_table)
        self.assertEqual(
            [column.label for column in fields["custom:10"].table_columns],
            ["조건", "결과"],
        )

    def test_workbook_uses_two_header_rows_and_item_row_blocks(self) -> None:
        workbook, summary = create_baseline_comparison_workbook(
            self.result,
            tracker_name="요구사항 (ID 20)",
            reference_label="현재 상태",
            comparison_label="R1 (2026-01-01)",
            selected_field_keys=self.selected_keys,
            generated_at=datetime(2026, 8, 6, 12, 0, 0),
        )
        sheet = workbook["비교 결과"]

        self.assertEqual(workbook.sheetnames, ["요약", "비교 결과"])
        self.assertEqual(sheet["A1"].value, "결과")
        self.assertIn("A1:A2", {str(item) for item in sheet.merged_cells.ranges})
        self.assertEqual(sheet["D1"].value, "기준 · 현재 상태")
        self.assertEqual(sheet["I1"].value, "비교 · R1 (2026-01-01)")
        self.assertEqual(sheet["G2"].value, "검증 표.조건")
        self.assertEqual(sheet["H2"].value, "검증 표.결과")
        self.assertIn("A3:A4", {str(item) for item in sheet.merged_cells.ranges})
        self.assertIn("F3:F4", {str(item) for item in sheet.merged_cells.ranges})
        self.assertEqual(sheet["G3"].value, "로그인")
        self.assertEqual(sheet["H3"].value, "예")
        self.assertEqual(sheet["G4"].value, "권한 확인")
        self.assertEqual(sheet["H4"].value, "예")
        self.assertEqual(sheet["M4"].value, "아니요")
        self.assertIn("• 홍길동 (ID 7)", sheet["F3"].value)
        self.assertIn("• 김영희 (ID 9)", sheet["F3"].value)
        self.assertEqual(sheet.freeze_panes, "D3")
        self.assertIsNone(sheet.auto_filter.ref)
        self.assertEqual(summary.item_count, 3)
        self.assertEqual(summary.data_row_count, 4)

    def test_exported_file_reopens_with_merged_ranges_and_without_formula_injection(self) -> None:
        injected = _item(4, name="=HYPERLINK('unsafe')", status="Open")
        result = compare_tracker_items(
            (),
            (injected,),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "baseline.xlsx"
            export_baseline_comparison_xlsx(
                result,
                path,
                tracker_name="요구사항",
                reference_label="현재 상태",
                comparison_label="R1",
                selected_field_keys=("name",),
            )
            workbook = load_workbook(path, data_only=False)

        self.assertEqual(workbook.sheetnames, ["요약", "비교 결과"])
        self.assertEqual(workbook["비교 결과"]["D3"].value, "'=HYPERLINK('unsafe')")

    def test_export_rejects_empty_selection(self) -> None:
        with self.assertRaises(BaselineExportError):
            create_baseline_comparison_workbook(
                self.result,
                tracker_name="요구사항",
                reference_label="현재 상태",
                comparison_label="R1",
                selected_field_keys=(),
            )


if __name__ == "__main__":
    unittest.main()
