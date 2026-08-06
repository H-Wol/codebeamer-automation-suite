from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter

from .tracker_baseline_compare import BaselineComparisonKind
from .tracker_baseline_compare import BaselineComparisonResult
from .tracker_baseline_compare import TrackerFieldDifference
from .tracker_baseline_compare import TrackerItemComparison
from .tracker_baseline_compare import TrackerTableColumn
from .tracker_baseline_compare import comparison_value_key
from .tracker_baseline_compare import table_field_rows


EXCEL_MAX_CELL_TEXT = 32767
EXCEL_MAX_COLUMNS = 16384
EXCEL_MAX_ROWS = 1048576


class BaselineExportError(RuntimeError):
    pass


@dataclass(frozen=True)
class BaselineExportField:
    field_key: str
    label: str
    is_table: bool = False
    table_columns: tuple[TrackerTableColumn, ...] = ()

    @property
    def column_count(self) -> int:
        return max(len(self.table_columns), 1)


@dataclass(frozen=True)
class BaselineExportSummary:
    item_count: int
    data_row_count: int
    selected_field_count: int
    output_path: str = ""


_KIND_LABELS = {
    BaselineComparisonKind.ADDED: "추가",
    BaselineComparisonKind.REMOVED: "삭제",
    BaselineComparisonKind.CHANGED: "변경",
    BaselineComparisonKind.UNCHANGED: "변경 없음",
}

_FONT_NAME = "Arial Unicode MS"
_GROUP_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_ADDED_FILL = PatternFill("solid", fgColor="DCFCE7")
_REMOVED_FILL = PatternFill("solid", fgColor="FEE2E2")
_CHANGED_FILL = PatternFill("solid", fgColor="FEF3C7")
_UNCHANGED_FILL = PatternFill("solid", fgColor="ECFDF5")
_WHITE_FONT = Font(name=_FONT_NAME, color="FFFFFF", bold=True)
_HEADER_FONT = Font(name=_FONT_NAME, color="1F2937", bold=True)
_BODY_FONT = Font(name=_FONT_NAME, color="1F2937")
_THIN_SIDE = Side(style="thin", color="CBD5E1")
_BORDER = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)
_TOP_WRAP = Alignment(vertical="top", wrap_text=True)
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)


def baseline_export_fields(
    result: BaselineComparisonResult,
) -> tuple[BaselineExportField, ...]:
    """전체 비교 결과의 필드를 처음 등장한 순서로 합친다."""
    fields: list[BaselineExportField] = []
    indexes: dict[str, int] = {}
    for item in result.items:
        for difference in item.fields:
            existing_index = indexes.get(difference.field_key)
            if existing_index is None:
                indexes[difference.field_key] = len(fields)
                fields.append(
                    BaselineExportField(
                        difference.field_key,
                        difference.label,
                        difference.is_table,
                        difference.table_columns,
                    )
                )
                continue
            existing = fields[existing_index]
            columns = _merge_columns(existing.table_columns, difference.table_columns)
            fields[existing_index] = BaselineExportField(
                existing.field_key,
                existing.label or difference.label,
                existing.is_table or difference.is_table,
                columns,
            )
    return tuple(fields)


def create_baseline_comparison_workbook(
    result: BaselineComparisonResult,
    *,
    tracker_name: str,
    reference_label: str,
    comparison_label: str,
    selected_field_keys: tuple[str, ...] | list[str],
    generated_at: datetime | None = None,
) -> tuple[Workbook, BaselineExportSummary]:
    fields_by_key = {field.field_key: field for field in baseline_export_fields(result)}
    selected_fields = tuple(
        fields_by_key[key]
        for key in dict.fromkeys(str(key) for key in selected_field_keys)
        if key in fields_by_key
    )
    if not selected_fields:
        raise BaselineExportError("내보낼 필드를 하나 이상 선택하세요.")

    flattened_count = sum(field.column_count for field in selected_fields)
    total_columns = 3 + flattened_count * 2
    if total_columns > EXCEL_MAX_COLUMNS:
        raise BaselineExportError(
            f"선택한 필드가 Excel 열 한도({EXCEL_MAX_COLUMNS:,}열)를 초과합니다."
        )

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "요약"
    comparison_sheet = workbook.create_sheet("비교 결과")
    generated = generated_at or datetime.now().astimezone()

    filtered_items = tuple(result.items)
    counts = _selected_counts(filtered_items, selected_fields)
    _populate_summary_sheet(
        summary_sheet,
        tracker_name=tracker_name,
        reference_label=reference_label,
        comparison_label=comparison_label,
        generated_at=generated,
        selected_fields=selected_fields,
        item_count=len(filtered_items),
        counts=counts,
    )
    data_row_count = _populate_comparison_sheet(
        comparison_sheet,
        items=filtered_items,
        selected_fields=selected_fields,
        reference_label=reference_label,
        comparison_label=comparison_label,
    )
    workbook.properties.title = "Codebeamer Baseline 비교 결과"
    workbook.properties.subject = str(tracker_name or "Tracker")
    workbook.properties.creator = "Codebeamer Automation Suite"
    return workbook, BaselineExportSummary(
        item_count=len(filtered_items),
        data_row_count=data_row_count,
        selected_field_count=len(selected_fields),
    )


def export_baseline_comparison_xlsx(
    result: BaselineComparisonResult,
    output_path: str | Path,
    *,
    tracker_name: str,
    reference_label: str,
    comparison_label: str,
    selected_field_keys: tuple[str, ...] | list[str],
) -> BaselineExportSummary:
    path = Path(output_path).expanduser()
    if path.suffix.casefold() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook, summary = create_baseline_comparison_workbook(
        result,
        tracker_name=tracker_name,
        reference_label=reference_label,
        comparison_label=comparison_label,
        selected_field_keys=selected_field_keys,
    )
    try:
        workbook.save(path)
    except Exception as exc:
        raise BaselineExportError("Excel 파일을 저장하지 못했습니다.") from exc
    return BaselineExportSummary(
        item_count=summary.item_count,
        data_row_count=summary.data_row_count,
        selected_field_count=summary.selected_field_count,
        output_path=str(path),
    )


def _merge_columns(
    *groups: tuple[TrackerTableColumn, ...],
) -> tuple[TrackerTableColumn, ...]:
    merged: list[TrackerTableColumn] = []
    seen: set[str] = set()
    for group in groups:
        for column in group:
            if column.column_key in seen:
                continue
            seen.add(column.column_key)
            merged.append(column)
    return tuple(merged)


def _selected_item_kind(
    item: TrackerItemComparison,
    selected_keys: set[str],
) -> BaselineComparisonKind:
    if item.kind in {BaselineComparisonKind.ADDED, BaselineComparisonKind.REMOVED}:
        return item.kind
    return (
        BaselineComparisonKind.CHANGED
        if any(
            field.field_key in selected_keys and field.is_changed
            for field in item.fields
        )
        else BaselineComparisonKind.UNCHANGED
    )


def _selected_counts(
    items: tuple[TrackerItemComparison, ...],
    fields: tuple[BaselineExportField, ...],
) -> dict[BaselineComparisonKind, int]:
    selected_keys = {field.field_key for field in fields}
    return {
        kind: sum(_selected_item_kind(item, selected_keys) == kind for item in items)
        for kind in BaselineComparisonKind
    }


def _populate_summary_sheet(
    sheet,
    *,
    tracker_name: str,
    reference_label: str,
    comparison_label: str,
    generated_at: datetime,
    selected_fields: tuple[BaselineExportField, ...],
    item_count: int,
    counts: dict[BaselineComparisonKind, int],
) -> None:
    sheet["A1"] = "Baseline 비교 내보내기"
    sheet["A1"].font = Font(
        name=_FONT_NAME,
        size=16,
        bold=True,
        color="FFFFFF",
    )
    sheet["A1"].fill = _GROUP_FILL
    sheet.merge_cells("A1:B1")
    rows = (
        ("트래커", tracker_name),
        ("기준", reference_label),
        ("비교", comparison_label),
        ("생성 시각", generated_at.isoformat(timespec="seconds")),
        ("선택 필드", "\n".join(field.label for field in selected_fields)),
        ("전체 아이템", item_count),
        ("추가", counts[BaselineComparisonKind.ADDED]),
        ("삭제", counts[BaselineComparisonKind.REMOVED]),
        ("변경", counts[BaselineComparisonKind.CHANGED]),
        ("변경 없음", counts[BaselineComparisonKind.UNCHANGED]),
    )
    for row_index, (label, value) in enumerate(rows, start=3):
        sheet.cell(row_index, 1, label)
        _set_safe_value(sheet.cell(row_index, 2), value, context=label)
        sheet.cell(row_index, 1).font = _HEADER_FONT
        sheet.cell(row_index, 1).fill = _HEADER_FILL
        for column in (1, 2):
            sheet.cell(row_index, column).border = _BORDER
            sheet.cell(row_index, column).alignment = _TOP_WRAP
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 48
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False


def _flattened_headers(
    fields: tuple[BaselineExportField, ...],
) -> tuple[tuple[BaselineExportField, TrackerTableColumn | None, str], ...]:
    headers: list[tuple[BaselineExportField, TrackerTableColumn | None, str]] = []
    for field in fields:
        if field.is_table and field.table_columns:
            for column in field.table_columns:
                headers.append((field, column, f"{field.label}.{column.label}"))
        else:
            headers.append((field, None, field.label))
    return tuple(headers)


def _populate_comparison_sheet(
    sheet,
    *,
    items: tuple[TrackerItemComparison, ...],
    selected_fields: tuple[BaselineExportField, ...],
    reference_label: str,
    comparison_label: str,
) -> int:
    headers = _flattened_headers(selected_fields)
    reference_start = 4
    reference_end = reference_start + len(headers) - 1
    comparison_start = reference_end + 1
    comparison_end = comparison_start + len(headers) - 1

    for column, label in enumerate(("결과", "아이템 ID", "아이템명"), start=1):
        sheet.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
        cell = sheet.cell(1, column, label)
        _style_group_header(cell)
    sheet.merge_cells(
        start_row=1,
        start_column=reference_start,
        end_row=1,
        end_column=reference_end,
    )
    sheet.merge_cells(
        start_row=1,
        start_column=comparison_start,
        end_row=1,
        end_column=comparison_end,
    )
    reference_group = sheet.cell(1, reference_start)
    comparison_group = sheet.cell(1, comparison_start)
    _set_safe_value(reference_group, f"기준 · {reference_label}", context="기준 제목")
    _set_safe_value(comparison_group, f"비교 · {comparison_label}", context="비교 제목")
    _style_group_header(reference_group)
    _style_group_header(comparison_group)

    for offset, (_field, _column, label) in enumerate(headers):
        for start in (reference_start, comparison_start):
            cell = sheet.cell(2, start + offset, label)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER_WRAP
            cell.border = _BORDER

    selected_keys = {field.field_key for field in selected_fields}
    row_index = 3
    for item in items:
        difference_by_key = {field.field_key: field for field in item.fields}
        row_count = _item_row_count(difference_by_key, selected_fields)
        if row_index + row_count - 1 > EXCEL_MAX_ROWS:
            raise BaselineExportError(
                f"아이템 #{item.item_id}에서 Excel 행 한도({EXCEL_MAX_ROWS:,}행)를 초과합니다."
            )
        start_row = row_index
        end_row = row_index + row_count - 1
        kind = _selected_item_kind(item, selected_keys)
        _merge_item_value(sheet, start_row, end_row, 1, _KIND_LABELS[kind])
        _merge_item_value(sheet, start_row, end_row, 2, item.item_id)
        _merge_item_value(sheet, start_row, end_row, 3, item.name)
        sheet.cell(start_row, 1).fill = _kind_fill(kind)
        sheet.cell(start_row, 1).font = Font(name=_FONT_NAME, bold=True)

        for offset, (export_field, table_column, _label) in enumerate(headers):
            difference = difference_by_key.get(export_field.field_key)
            reference_column = reference_start + offset
            comparison_column = comparison_start + offset
            if difference is None:
                _merge_item_value(sheet, start_row, end_row, reference_column, "-")
                _merge_item_value(sheet, start_row, end_row, comparison_column, "-")
                continue
            if not export_field.is_table or table_column is None:
                _merge_item_value(
                    sheet,
                    start_row,
                    end_row,
                    reference_column,
                    difference.reference_text(),
                    context=f"#{item.item_id} {difference.label} 기준",
                )
                _merge_item_value(
                    sheet,
                    start_row,
                    end_row,
                    comparison_column,
                    difference.comparison_text(),
                    context=f"#{item.item_id} {difference.label} 비교",
                )
                _style_scalar_difference(
                    sheet.cell(start_row, reference_column),
                    sheet.cell(start_row, comparison_column),
                    difference,
                )
                continue
            _write_table_column(
                sheet,
                start_row=start_row,
                row_count=row_count,
                reference_column=reference_column,
                comparison_column=comparison_column,
                item_id=item.item_id,
                difference=difference,
                table_column=table_column,
            )

        for row in range(start_row, end_row + 1):
            for column in range(1, comparison_end + 1):
                cell = sheet.cell(row, column)
                if not isinstance(cell, Cell):
                    continue
                cell.border = _BORDER
                cell.alignment = _TOP_WRAP
        row_index = end_row + 1

    sheet.freeze_panes = "D3"
    widths = {1: 13, 2: 13, 3: 32}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(reference_start, comparison_end + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 24
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 36
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:2"
    return max(row_index - 3, 0)


def _style_group_header(cell: Cell) -> None:
    cell.fill = _GROUP_FILL
    cell.font = _WHITE_FONT
    cell.alignment = _CENTER_WRAP
    cell.border = _BORDER


def _item_row_count(
    difference_by_key: dict[str, TrackerFieldDifference],
    selected_fields: tuple[BaselineExportField, ...],
) -> int:
    row_count = 1
    for field in selected_fields:
        if not field.is_table:
            continue
        difference = difference_by_key.get(field.field_key)
        if difference is None:
            continue
        row_count = max(
            row_count,
            len(table_field_rows(difference.reference)),
            len(table_field_rows(difference.comparison)),
        )
    return row_count


def _merge_item_value(
    sheet,
    start_row: int,
    end_row: int,
    column: int,
    value: Any,
    *,
    context: str = "셀",
) -> None:
    if end_row > start_row:
        sheet.merge_cells(
            start_row=start_row,
            start_column=column,
            end_row=end_row,
            end_column=column,
        )
    cell = sheet.cell(start_row, column)
    _set_safe_value(cell, value, context=context)
    cell.alignment = _TOP_WRAP
    cell.border = _BORDER


def _write_table_column(
    sheet,
    *,
    start_row: int,
    row_count: int,
    reference_column: int,
    comparison_column: int,
    item_id: int,
    difference: TrackerFieldDifference,
    table_column: TrackerTableColumn,
) -> None:
    reference_rows = table_field_rows(difference.reference)
    comparison_rows = table_field_rows(difference.comparison)
    missing = object()
    for offset in range(row_count):
        reference_value = (
            reference_rows[offset].get(table_column.column_key, missing)
            if offset < len(reference_rows)
            else missing
        )
        comparison_value = (
            comparison_rows[offset].get(table_column.column_key, missing)
            if offset < len(comparison_rows)
            else missing
        )
        reference_cell = sheet.cell(start_row + offset, reference_column)
        comparison_cell = sheet.cell(start_row + offset, comparison_column)
        _set_safe_value(
            reference_cell,
            "" if reference_value is missing else _display_cell_value(reference_value),
            context=f"#{item_id} {difference.label}.{table_column.label} 기준",
        )
        _set_safe_value(
            comparison_cell,
            "" if comparison_value is missing else _display_cell_value(comparison_value),
            context=f"#{item_id} {difference.label}.{table_column.label} 비교",
        )
        reference_cell.alignment = _TOP_WRAP
        comparison_cell.alignment = _TOP_WRAP
        if reference_value is missing and comparison_value is missing:
            continue
        if comparison_value is missing:
            reference_cell.fill = _ADDED_FILL
        elif reference_value is missing:
            comparison_cell.fill = _REMOVED_FILL
        elif comparison_value_key(reference_value) != comparison_value_key(comparison_value):
            reference_cell.fill = _CHANGED_FILL
            comparison_cell.fill = _CHANGED_FILL


def _display_cell_value(value: Any) -> str:
    if value is None or value == "" or value == [] or value == {}:
        return ""
    if isinstance(value, bool):
        return "예" if value else "아니요"
    if isinstance(value, dict):
        name = value.get("name") or value.get("summary") or value.get("label")
        reference_id = value.get("id")
        if name not in (None, "") and reference_id not in (None, ""):
            return f"{name} (ID {reference_id})"
        if name not in (None, ""):
            return str(name)
        if reference_id not in (None, ""):
            return f"ID {reference_id}"
        return "\n".join(
            f"{key}: {_display_cell_value(nested)}" for key, nested in value.items()
        )
    if isinstance(value, (list, tuple)):
        return "\n".join(f"• {_display_cell_value(item)}" for item in value)
    return str(value)


def _style_scalar_difference(
    reference_cell: Cell,
    comparison_cell: Cell,
    difference: TrackerFieldDifference,
) -> None:
    if not difference.is_changed:
        return
    if _is_empty(difference.comparison) and not _is_empty(difference.reference):
        reference_cell.fill = _ADDED_FILL
    elif _is_empty(difference.reference) and not _is_empty(difference.comparison):
        comparison_cell.fill = _REMOVED_FILL
    else:
        reference_cell.fill = _CHANGED_FILL
        comparison_cell.fill = _CHANGED_FILL


def _is_empty(value: Any) -> bool:
    return value is None or value == "" or value == [] or value == {}


def _kind_fill(kind: BaselineComparisonKind) -> PatternFill:
    return {
        BaselineComparisonKind.ADDED: _ADDED_FILL,
        BaselineComparisonKind.REMOVED: _REMOVED_FILL,
        BaselineComparisonKind.CHANGED: _CHANGED_FILL,
        BaselineComparisonKind.UNCHANGED: _UNCHANGED_FILL,
    }[kind]


def _set_safe_value(cell: Cell, value: Any, *, context: str) -> None:
    cell.font = _BODY_FONT
    if isinstance(value, str):
        normalized = value.replace("\x00", "")
        if len(normalized) > EXCEL_MAX_CELL_TEXT:
            raise BaselineExportError(
                f"{context} 내용이 Excel 셀 길이 한도({EXCEL_MAX_CELL_TEXT:,}자)를 초과합니다."
            )
        if normalized.startswith(("=", "+", "@")) or (
            normalized.startswith("-") and normalized != "-"
        ):
            normalized = f"'{normalized}"
        cell.value = normalized
        return
    cell.value = value


__all__ = [
    "BaselineExportError",
    "BaselineExportField",
    "BaselineExportSummary",
    "baseline_export_fields",
    "create_baseline_comparison_workbook",
    "export_baseline_comparison_xlsx",
]
